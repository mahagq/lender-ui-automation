"""
create_test_sheet_xlsx.py
Generates CP_UI_Automation_Test_Cases.xlsx with 5 sheets:
  Sheet 1: Login Left Component   (9 TCs)
  Sheet 2: Login                  (21 TCs)
  Sheet 3: OTP Verification       (14 TCs, 3 skipped)
  Sheet 4: Bypass Authentication  (4 TCs, 2 skipped)
  Sheet 5: Dashboard              (33 TCs, 2 skipped)

Requirements:
    pip install openpyxl

Run:
    python3 create_test_sheet_xlsx.py
"""

from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

OUTPUT_FILE = "CP_UI_Automation_Test_Cases.xlsx"

HEADERS = [
    "TC ID", "Test Case Name", "Suite", "Tags",
    "Type", "Description", "Expected Result",
    "Test Mobile / Data", "Status", "Notes"
]

# Column widths (approximate characters)
COL_WIDTHS = [10, 36, 20, 28, 12, 45, 36, 20, 9, 30]

# ── FILL COLORS ───────────────────────────────────────────────────────────────
def hex_fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)

FILL_HEADER   = hex_fill("D9D9D9")   # grey
FILL_PASS     = hex_fill("B6D7A8")   # green
FILL_SKIP     = hex_fill("FFE599")   # yellow
FILL_FAIL     = hex_fill("EA9999")   # red
FILL_WHITE    = hex_fill("FFFFFF")

# Tab colors (hex for sheet tab — openpyxl uses ThemeColor or RGB)
TAB_COLORS = {
    "Login Left Component":  "AED6E6",  # blue
    "Login":                 "BFB9E6",  # purple
    "OTP Verification":      "FFD699",  # orange
    "Bypass Authentication": "F4B8D1",  # pink
    "Dashboard":             "B6E6CC",  # green
}

STATUS_FILL = {"PASS": FILL_PASS, "SKIP": FILL_SKIP, "FAIL": FILL_FAIL}

# ── BORDERS ───────────────────────────────────────────────────────────────────
thin  = Side(style="thin",  color="999999")
light = Side(style="thin",  color="CCCCCC")
outer_border = Border(left=thin, right=thin, top=thin, bottom=thin)
inner_border = Border(left=light, right=light, top=light, bottom=light)

# ── FONTS ─────────────────────────────────────────────────────────────────────
font_header = Font(name="Calibri", bold=True, size=10, color="333333")
font_id     = Font(name="Calibri", bold=True, size=9)
font_status = Font(name="Calibri", bold=True, size=9)
font_data   = Font(name="Calibri", size=9)

# ── ALIGNMENT ─────────────────────────────────────────────────────────────────
align_center  = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_top     = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
align_tc_id   = Alignment(horizontal="center", vertical="top",    wrap_text=False)
align_status  = Alignment(horizontal="center", vertical="top",    wrap_text=False)

# ── TEST DATA ─────────────────────────────────────────────────────────────────

LEFT_COMPONENT = [
    ["TC-101", "GQ Logo Is Displayed In Left Section",
     "Login Left Component", "left-component, logo, ui",
     "UI", "The GrayQuest SVG logo should be present and visible at the top of the left panel.",
     "SVG logo is visible", "", "PASS", ""],
    ["TC-102", "Register For EMI Plans Heading Text Is Correct",
     "Login Left Component", "left-component, ui, heading",
     "Content", "Left panel main heading split across two elements: 'Register for' + 'EMI Plans'.",
     "'Register for' and 'EMI Plans' text verified", "", "PASS", ""],
    ["TC-103", "Subheading Paragraph Text Is Correct",
     "Login Left Component", "left-component, ui, text",
     "Content", "data-testid='info-paragraph' should display the subheading paragraph exactly.",
     "Exact paragraph text matches source", "", "PASS", ""],
    ["TC-104", "Trusted By Image Is Visible",
     "Login Left Component", "left-component, ui, image",
     "UI", "The 'Trusted by 5M+ Parents' image section should be visible in the left panel.",
     "Trusted by image is visible", "", "PASS", ""],
    ["TC-105", "About Us Link Opens About Us Page",
     "Login Left Component", "left-component, navigation, new-tab",
     "Navigation", "About Us link in footer should open About Us page in a new tab.",
     "New tab title contains 'About Us'", "", "PASS", ""],
    ["TC-106", "Privacy Policy Link Opens Privacy Policy Page",
     "Login Left Component", "left-component, navigation, new-tab",
     "Navigation", "Privacy Policy link in footer should open Privacy Policy page in a new tab.",
     "New tab title contains 'Privacy Policy'", "", "PASS", ""],
    ["TC-107", "Terms And Conditions Link Opens Terms And Conditions Page",
     "Login Left Component", "left-component, navigation, new-tab",
     "Navigation", "T&C link in footer should open T&C page in a new tab.",
     "New tab title contains 'Terms & Conditions'", "", "PASS", ""],
    ["TC-108", "Contact Us Link Opens Contact Us Page",
     "Login Left Component", "left-component, navigation, new-tab",
     "Navigation", "Contact Us link in footer should open Contact Us page in a new tab.",
     "New tab title contains 'Contact Us'", "", "PASS", ""],
    ["TC-109", "Copyright Footer Text Is Correct",
     "Login Left Component", "left-component, ui, footer-text",
     "Content", "Copyright notice at bottom of left panel should display current year and company name.",
     "'© <year> GrayQuest Education Finance Pvt. Ltd.' matches", "", "PASS", ""],
]

LOGIN = [
    ["TC-201", "Login Page URL And Title Match",
     "Login", "login, smoke, url",
     "Smoke", "Navigate to base URL and verify page URL and browser tab title.",
     "URL = BASE_URL/ and title contains 'GrayQuest Education Finance'", "", "PASS",
     "Asserts title contains 'GrayQuest Education Finance' (Should Contain — resilient to title prefix changes)"],
    ["TC-207", "Mobile Number Input Field Is Visible And Editable",
     "Login", "login, ui, field",
     "UI", "Mobile number input field should be visible and accept user input.",
     "Input is visible and enabled", "", "PASS", ""],
    ["TC-208", "Login With Empty Mobile Number",
     "Login", "login, validation, negative",
     "Validation", "Clicking Send OTP without entering a mobile number triggers validation error.",
     "'Please enter your mobile number'", "", "PASS", ""],
    ["TC-209", "Login With Spaces In Mobile Number",
     "Login", "login, validation, negative",
     "Validation", "Entering only whitespace is treated as empty — numeric input filters spaces.",
     "'Please enter your mobile number'", "10 spaces", "PASS", ""],
    ["TC-210", "Login With Alphabets In Mobile Number",
     "Login", "login, validation, negative",
     "Validation", "Purely alphabetic string triggers mobile regex validation.",
     "'Please enter valid mobile number'", "abcdefghij", "PASS", ""],
    ["TC-211", "Login With Invalid Mobile Number Format",
     "Login", "login, validation, negative",
     "Validation", "Mobile number containing letters triggers Yup Regex.mobille_number validation.",
     "'Please enter valid mobile number'", "123abc4567", "PASS", ""],
    ["TC-212", "Login With Alphanumeric Mobile Number",
     "Login", "login, validation, negative",
     "Validation", "Mix of letters and digits triggers mobile regex validation.",
     "'Please enter valid mobile number'", "abc1234567", "PASS", ""],
    ["TC-213", "Login With Special Characters In Mobile Number",
     "Login", "login, validation, negative",
     "Validation", "Special characters trigger mobile regex validation.",
     "'Please enter valid mobile number'", "@#$%^&*()!", "PASS", ""],
    ["TC-214", "Login With Short Mobile Number",
     "Login", "login, validation, negative",
     "Validation", "Fewer than 10 digits — Yup evaluates regex before length, regex fires first.",
     "'Please enter valid mobile number'", "12345", "PASS", ""],
    ["TC-215", "Mobile Number Field Enforces 10 Digit Max Length",
     "Login", "login, validation, field",
     "Validation", "Mobile input has maxLength=10. Entering 14 digits results in only 10 stored.",
     "Field value = first 10 digits", "12345678901234", "PASS", ""],
    ["TC-216", "Valid Login With Correct Mobile Number",
     "Login", "login, valid, smoke",
     "Smoke", "Valid 10-digit mobile + consent checkbox + Send OTP → redirect to /verify.",
     "Browser navigates to /verify, OTP form visible", "9876543210", "PASS", ""],
    ["TC-217", "Valid Mobile Starting With 6 To 9 Navigates To OTP Page",
     "Login", "login, valid, otp-redirect",
     "Positive", "Valid 10-digit mobile starting 6–9 with consent → OTP page title 'Enter 4 Digit OTP'.",
     "OTP page title = 'Enter 4 Digit OTP'", "9876543215", "PASS", ""],
    ["TC-218", "Login Without Consent",
     "Login", "login, validation, negative",
     "Validation", "Valid mobile but NO consent checkbox → Yup bool().oneOf([true]) validation.",
     "'Please accept Terms and Conditions'", "9876543210", "PASS", ""],
    ["TC-219", "Send OTP With Consent Checked Navigates To OTP Page",
     "Login", "login, valid, otp-redirect",
     "Positive", "Valid mobile + consent checked → navigate to /verify, title = 'Enter 4 Digit OTP'.",
     "OTP page title = 'Enter 4 Digit OTP'", "9876543216", "PASS", ""],
    ["TC-220", "Send OTP Button Is Enabled With Mobile And Consent",
     "Login", "login, validation, button-state",
     "UI State", "Valid mobile + consent checkbox checked → Send OTP button is enabled.",
     "Send OTP button state = enabled", "9876543210", "PASS", ""],
    ["TC-221", "Login Form Title Is Visible",
     "Login", "login, ui, title",
     "UI", "Login form should display the title 'What is your Mobile Number?' via data-testid='desktop-view-form-title'.",
     "Title = 'What is your Mobile Number?'", "", "PASS", ""],
    ["TC-222", "Mobile Input Note Text Is Visible",
     "Login", "login, ui, note",
     "UI", "When no validation error, note below mobile input: 'Please enter the mobile number which you generally use for your banking purposes.'",
     "data-testid='mobile-number-input-note' visible and contains note text", "", "PASS", ""],
    ["TC-223", "Terms Of Service Link In Consent Opens New Tab",
     "Login", "login, ui, consent, navigation",
     "Navigation", "The 'Terms of Services' link inside the consent label should open the T&C page in a new tab.",
     "New tab title contains 'Terms'", "", "PASS", ""],
    ["TC-224", "Privacy Policy Link In Consent Opens New Tab",
     "Login", "login, ui, consent, navigation",
     "Navigation", "The 'Privacy Policy' link inside the consent label should open the Privacy Policy page in a new tab.",
     "New tab title contains 'Privacy'", "", "PASS", ""],
    ["TC-225", "Mobile Number Prefill Via Query Param Navigates To OTP Page",
     "Login", "login, valid, query-param",
     "Positive", "Navigating to /?m=<mobile> pre-fills the mobile field, auto-checks consent, auto-submits, and redirects to /verify.",
     "URL contains /verify", "9876543227", "PASS", ""],
    ["TC-226", "Mobile Footer Is Present On Login Page",
     "Login", "login, ui, mobile-footer",
     "UI", "MobileFooter component (data-testid='mobile-footer-container') should be present in the DOM on the login page.",
     "data-testid='mobile-footer-container' attached in DOM", "", "PASS",
     "Asserts 'attached' state — element is CSS-hidden on desktop viewport"],
]

OTP = [
    ["TC-301", "OTP Input Field Is Visible And Editable",
     "OTP Verification", "otp, ui, field",
     "UI", "On the OTP page, the OTP input field should be visible and accept user input.",
     "OTP input is visible and enabled", "9876543218", "PASS", ""],
    ["TC-302", "OTP Verification With Empty OTP",
     "OTP Verification", "otp, validation, negative",
     "Validation", "Clicking Verify OTP without entering any value triggers React Hook Form validation.",
     "'Please enter 4 digit OTP'", "9876543211", "PASS", ""],
    ["TC-303", "OTP Verification With Short OTP",
     "OTP Verification", "otp, validation, negative",
     "Validation", "Entering fewer than 4 digits triggers Yup .length(4) validation.",
     "'OTP number must contain 4 digits'", "9876543212 / OTP: 123", "PASS", ""],
    ["TC-304", "Verify OTP Button Is Enabled After OTP Entry",
     "OTP Verification", "otp, ui, button-state",
     "UI State", "After entering a 4-digit OTP, the Verify OTP button should transition to enabled.",
     "Verify OTP button state = enabled", "9876543225", "PASS", ""],
    ["TC-305", "Valid OTP Redirects To Dashboard With Greeting",
     "OTP Verification", "otp, valid, smoke, dashboard",
     "Smoke", "Correct UAT OTP (9999) → verifyOtpPost → redirect to /my-applications with 'Hello,' greeting.",
     "URL = /my-applications, 'Hello,' visible", "9876543220 / OTP: 9999", "PASS",
     "Three-state rate limit handler: polls redirect, then checks dashboard/OTP/login page after 45s"],
    ["TC-306", "Wrong OTP Shows Incorrect OTP Toast",
     "OTP Verification", "otp, negative, toast",
     "Negative", "Wrong but structurally valid OTP → verifyOtpPost error. User stays on /verify.",
     "URL stays on /verify (primary); 'Incorrect OTP' toast shown (secondary, non-blocking)",
     "9876543219 / OTP: 1234", "PASS",
     "Toast non-blocking — verifyOtpPost may return 429 in UAT if rate-limited"],
    ["TC-307", "Back Button Returns To Login Page",
     "OTP Verification", "otp, navigation, back-button",
     "Navigation", "Clicking back button on OTP page returns to login page.",
     "Login form visible, title = 'What is your Mobile Number?'", "9876543217", "PASS", ""],
    ["TC-308", "Back Button Clears Mobile Number Input",
     "OTP Verification", "otp, navigation, back-button",
     "Navigation", "After back button click, mobile input field should be cleared (empty).",
     "Mobile input value = empty", "9876543222", "PASS", ""],
    ["TC-309", "Entered Mobile Number Is Displayed On OTP Page",
     "OTP Verification", "otp, ui, mobile-display",
     "UI", "Mobile number entered on login form should be displayed on OTP page.",
     "data-testid='mobile-number-info' contains submitted mobile", "9876543221", "PASS", ""],
    ["TC-310", "OTP Countdown Timer Shows Resend Text",
     "OTP Verification", "otp, ui, timer",
     "UI", "Immediately after OTP page loads, timer container visible with countdown text.",
     "Timer contains 'Did not receive OTP? Resend in'", "9876543223", "PASS", ""],
    ["TC-314", "OTP With Non-Numeric Characters Shows Error",
     "OTP Verification", "otp, validation, negative",
     "Validation", "Entering non-numeric characters (e.g. 'abcd') into the OTP field triggers Yup .matches(Regex.otp) validation.",
     "'Only numeric values allowed'", "9876543226 / OTP: abcd", "PASS",
     "Input has inputMode='numeric' but type='text' — desktop keyboards can type letters"],
    ["TC-311", "OTP Not Received",
     "OTP Verification", "otp, manual, robot:skip",
     "Manual", "Manual test — verify UX when OTP is not received (carrier issues).",
     "N/A — manual verification", "", "SKIP",
     "Cannot automate — requires external SMS gateway dependency"],
    ["TC-312", "OTP Timer Disabled State",
     "OTP Verification", "otp, manual, robot:skip",
     "Manual", "Manual test — verify timer appearance before countdown starts.",
     "N/A — manual verification", "", "SKIP",
     "Cannot reliably automate within normal timing constraints"],
    ["TC-313", "Resend OTP After Timer Expires",
     "OTP Verification", "otp, slow, resend, robot:skip",
     "Slow / Manual", "After 180s countdown, Resend OTP link becomes visible. Clicking resets timer.",
     "Resend link visible after 180s; timer resets after click", "", "SKIP",
     "~3 min wait. Run with: robot --include resend tests/otp_tests.robot"],
]

BYPASS_AUTH = [
    ["TC-401", "Bypass Auth With Valid Mobile Redirects Away From Instant Eligibility",
     "Bypass Authentication", "bypass-auth, smoke, redirect",
     "Smoke", "Navigating to /instant-eligibility?m=<valid_mobile> authenticates via bypassAuthPost and redirects away from /instant-eligibility.",
     "URL no longer contains /instant-eligibility", "9876543210", "PASS", ""],
    ["TC-402", "Bypass Auth Without Query Param Redirects To Login Page",
     "Bypass Authentication", "bypass-auth, redirect, negative",
     "Negative", "Navigating to /instant-eligibility without ?m= redirects to / (login page) with login form visible.",
     "Login form visible; URL does not contain /instant-eligibility", "", "PASS", ""],
    ["TC-403", "Bypass Auth No Account Error State Is Visible",
     "Bypass Authentication", "bypass-auth, negative, robot:skip",
     "Skipped", "When bypassAuthPost fails (mobile has no ERP account), no-account error state should show image and message.",
     "No-account image and message visible", "", "SKIP",
     "Requires UAT mobile that fails bypassAuthPost — not yet identified"],
    ["TC-404", "Sign In Button On No Account State Navigates To Login",
     "Bypass Authentication", "bypass-auth, negative, navigation, robot:skip",
     "Skipped", "On no-account error state, clicking Sign In should navigate to / (login page).",
     "Login form visible after Sign In click", "", "SKIP",
     "Depends on TC-403 precondition — bypassAuthPost failure state"],
]

DASHBOARD = [
    ["TC-4001", "EMI Applications Are Visible On Dashboard",
     "Dashboard", "dashboard, emi, smoke",
     "Smoke", "Dashboard EMI tab should display at least one EMI card after login.",
     "At least one EMI card visible", "8108556708", "PASS",
     "PRECONDITION: account must have active EMI applications"],
    ["TC-4002", "Auto Debit Applications On Dashboard",
     "Dashboard", "dashboard, auto-debit, tab-navigation",
     "Navigation", "The Auto Payments (AD) tab navigation button should be visible on the dashboard.",
     "data-testid='ad-tab-navigation-button' visible", "8108556708", "PASS", ""],
    ["TC-4003", "Payment Gateway Applications Are Visible",
     "Dashboard", "dashboard, pg, smoke",
     "Smoke", "Transaction History tab should display Online and Offline sub-tab navigation.",
     "Online and Offline sub-tabs visible", "8108556708", "PASS", ""],
    ["TC-4004", "Stages And CTA Of Stages",
     "Dashboard", "dashboard, stages, robot:skip",
     "Skipped", "Validates stages and call-to-action elements on the dashboard.",
     "N/A", "", "SKIP",
     "Empty body in Java suite — no data-testid for stage elements"],
    ["TC-4005", "Back Button Is Not Present On Dashboard",
     "Dashboard", "dashboard, navigation, robot:skip",
     "Skipped", "Verifies back button behaviour on the dashboard.",
     "N/A", "", "SKIP",
     "Empty body in Java suite — no back button on dashboard"],
    ["TC-4006", "GrayQuest Logo Is Visible On Dashboard",
     "Dashboard", "dashboard, logo, ui, smoke",
     "Smoke", "GQ desktop logo visible. Also asserts title contains 'Hello,' and subtitle contains 'Autopay subscriptions'.",
     "Logo visible; title = 'Hello,'; subtitle contains 'Autopay subscriptions'", "8108556708", "PASS", ""],
    ["TC-4007", "Profile Icon Shows FAQ And Logout Options",
     "Dashboard", "dashboard, profile, ui",
     "UI", "Clicking profile icon reveals dropdown with FAQ and Logout options.",
     "profile-options-container visible; FAQ and Logout options visible", "8108556708", "PASS", ""],
    ["TC-4008", "FAQ Link Navigates To FAQ Page",
     "Dashboard", "dashboard, profile, navigation",
     "Navigation", "Clicking FAQ in profile dropdown navigates to FAQ page.",
     "URL contains /faq", "8108556708", "PASS",
     "Asserted via URL only — FAQ h1 has no data-testid (TN3 resolution)"],
    ["TC-4032", "Insurance Tab Button Is Visible On Dashboard",
     "Dashboard", "dashboard, insurance, tab-navigation",
     "Navigation", "The Insurance tab navigation button (data-testid='ins-tab-navigation-button') should be visible on the dashboard.",
     "Insurance tab button visible", "8108556708", "PASS", ""],
    ["TC-4033", "Insurance Tab Click Updates URL Hash",
     "Dashboard", "dashboard, insurance, tab-navigation",
     "Navigation", "Clicking the Insurance tab should update the URL hash to #ins and load the Insurance content area.",
     "URL contains #ins", "8108556708", "PASS",
     "InsuranceTab has no data-testid attributes — URL hash is the primary assertion"],
    ["TC-4009", "Logout Returns To Login Page",
     "Dashboard", "dashboard, profile, logout, session",
     "Session", "Logout clears auth session and redirects to login page. MUST run last.",
     "Login form visible after logout", "8108556708", "PASS",
     "TC-4009 MUST be last test — destroys shared session (TN1 resolution)"],
    ["TC-4010", "Create New Application Button Has Visible Effect",
     "Dashboard", "dashboard, application, smoke",
     "Smoke", "Button click produces visible response: modal OR URL change from /my-applications.",
     "Modal visible OR URL changed", "8108556708", "PASS",
     "WARN expected in UAT (no modal — inline form). TN5 dual-branch per O2."],
    ["TC-4011", "EMI Tab Click Shows EMI Content",
     "Dashboard", "dashboard, emi, tab-navigation",
     "Navigation", "Clicking EMI tab updates URL hash to #emi and shows at least one EMI card.",
     "URL contains #emi; EMI card visible", "8108556708", "PASS",
     "PRECONDITION: account must have active EMI applications"],
    ["TC-4012", "Auto Payments Tab Click",
     "Dashboard", "dashboard, auto-payments, tab-navigation",
     "Navigation", "Clicking the Auto Payments tab button should update the URL hash to #ad and load the Auto Payments content area.",
     "URL contains #ad", "8108556708", "PASS", ""],
    ["TC-4013", "Transaction History Tab Click",
     "Dashboard", "dashboard, pg, tab-navigation, smoke",
     "Smoke", "Clicking Transaction History tab updates URL hash to #direct and shows sub-tabs.",
     "URL contains #direct; Online and Offline sub-tabs visible", "8108556708", "PASS", ""],
    ["TC-4014", "Offline Payments Sub-tab Is Accessible",
     "Dashboard", "dashboard, pg, sub-tab",
     "Navigation", "Offline Payments sub-tab should be visible and clickable.",
     "Offline sub-tab clickable", "8108556708", "PASS", ""],
    ["TC-4015", "Online Payments Sub-tab Is Accessible",
     "Dashboard", "dashboard, pg, sub-tab",
     "Navigation", "Online Payments sub-tab should be visible and clickable.",
     "Online sub-tab clickable", "8108556708", "PASS", ""],
    ["TC-4016", "Filter Panel Is Visible For Online Payments",
     "Dashboard", "dashboard, pg, filter, online",
     "UI", "Filter button on Online sub-tab opens filter panel with all time and status options.",
     "All time + status filter options visible", "8108556708", "PASS", ""],
    ["TC-4017", "Filter By All Months For Online Payments",
     "Dashboard", "dashboard, pg, filter, online, time-filter",
     "Filter", "'All Months' time filter option is selectable on Online sub-tab.",
     "'All Months' filter selectable", "8108556708", "PASS", ""],
    ["TC-4018", "Filter By This Month For Online Payments",
     "Dashboard", "dashboard, pg, filter, online, time-filter",
     "Filter", "'This Month' time filter option is selectable on Online sub-tab.",
     "'This Month' filter selectable", "8108556708", "PASS", ""],
    ["TC-4019", "Filter By This Year For Online Payments",
     "Dashboard", "dashboard, pg, filter, online, time-filter",
     "Filter", "'This Year' time filter option is selectable on Online sub-tab.",
     "'This Year' filter selectable", "8108556708", "PASS", ""],
    ["TC-4020", "Filter By Last Year For Online Payments",
     "Dashboard", "dashboard, pg, filter, online, time-filter",
     "Filter", "'Last Year' time filter option is selectable on Online sub-tab.",
     "'Last Year' filter selectable", "8108556708", "PASS", ""],
    ["TC-4021", "Filter By Status Paid For Online Payments",
     "Dashboard", "dashboard, pg, filter, online, status-filter",
     "Filter", "'Paid' status filter option is selectable on Online sub-tab.",
     "'Paid' filter selectable", "8108556708", "PASS", ""],
    ["TC-4022", "Filter By Status Failed For Online Payments",
     "Dashboard", "dashboard, pg, filter, online, status-filter",
     "Filter", "'Failed' status filter option is selectable on Online sub-tab.",
     "'Failed' filter selectable", "8108556708", "PASS", ""],
    ["TC-4023", "Filter Panel Is Visible For Offline Payments",
     "Dashboard", "dashboard, pg, filter, offline",
     "UI", "Filter button on Offline sub-tab opens filter panel with time filter options only.",
     "Time filter options visible (no status filters on Offline)", "8108556708", "PASS", ""],
    ["TC-4024", "Filter By All Months For Offline Payments",
     "Dashboard", "dashboard, pg, filter, offline, time-filter",
     "Filter", "'All Months' time filter is selectable on Offline sub-tab.",
     "'All Months' filter selectable", "8108556708", "PASS", ""],
    ["TC-4025", "Filter By This Month For Offline Payments",
     "Dashboard", "dashboard, pg, filter, offline, time-filter",
     "Filter", "'This Month' time filter option should be selectable on the Offline sub-tab.",
     "'This Month' filter selectable", "8108556708", "PASS", ""],
    ["TC-4026", "Filter By This Year For Offline Payments",
     "Dashboard", "dashboard, pg, filter, offline, time-filter",
     "Filter", "'This Year' time filter option should be selectable on the Offline sub-tab.",
     "'This Year' filter selectable", "8108556708", "PASS", ""],
    ["TC-4027", "Filter By Last Year For Offline Payments",
     "Dashboard", "dashboard, pg, filter, offline, time-filter",
     "Filter", "'Last Year' time filter is selectable on Offline sub-tab.",
     "'Last Year' filter selectable", "8108556708", "PASS", ""],
    ["TC-4028", "Clear Filter Button Works",
     "Dashboard", "dashboard, pg, filter, clear-filter",
     "Filter", "After selecting a filter, clear filter button resets the selection.",
     "Clear filter button clickable; selection resets", "8108556708", "PASS", ""],
    ["TC-4029", "Apply Filter Button Works",
     "Dashboard", "dashboard, pg, filter, apply-filter",
     "Filter", "After selecting a filter, apply filter button applies the selection.",
     "Apply filter button clickable", "8108556708", "PASS", ""],
    ["TC-4030", "Download Icon Is Visible For Offline Payment",
     "Dashboard", "dashboard, pg, download, offline",
     "Download", "Download button for first transaction row on Offline sub-tab should be visible.",
     "Download button visible on first Offline row", "7757855472 (PG_MOBILE)", "PASS",
     "Switches to PG_MOBILE via Switch To Account. TC-4031 inherits session."],
    ["TC-4031", "Download Icon Is Visible For Online Payment",
     "Dashboard", "dashboard, pg, download, online",
     "Download", "Download button for first transaction row on Online sub-tab should be visible.",
     "Download button visible on first Online row", "7757855472 (PG_MOBILE)", "PASS",
     "Inherits PG_MOBILE session from TC-4030. Canonical from DashboardTestClass3."],
]

SHEETS = [
    {"name": "Login Left Component",  "data": LEFT_COMPONENT},
    {"name": "Login",                 "data": LOGIN},
    {"name": "OTP Verification",      "data": OTP},
    {"name": "Bypass Authentication", "data": BYPASS_AUTH},
    {"name": "Dashboard",             "data": DASHBOARD},
]


def write_sheet(wb, sheet_def):
    name = sheet_def["name"]
    rows = sheet_def["data"]

    ws = wb.create_sheet(title=name)
    ws.sheet_properties.tabColor = TAB_COLORS[name]

    # ── Header row ────────────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill   = FILL_HEADER
        cell.font   = font_header
        cell.alignment = align_center
        cell.border = outer_border

    ws.freeze_panes = "A2"

    # ── Column widths ─────────────────────────────────────────────────────────
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, row in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 55
        status = row[8]
        row_fill = STATUS_FILL.get(status, FILL_WHITE)

        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill   = row_fill
            cell.border = inner_border

            if col_idx == 1:   # TC ID
                cell.font      = font_id
                cell.alignment = align_tc_id
            elif col_idx == 9:  # Status
                cell.font      = font_status
                cell.alignment = align_status
            else:
                cell.font      = font_data
                cell.alignment = align_top

    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    for sheet_def in SHEETS:
        write_sheet(wb, sheet_def)

    wb.save(OUTPUT_FILE)

    # ── Summary ───────────────────────────────────────────────────────────────
    print(f"\nSaved: {OUTPUT_FILE}\n")
    print(f"{'Sheet':<25}  {'TCs':>4}  PASS  SKIP  FAIL")
    print("─" * 50)
    total = 0
    for s in SHEETS:
        d = s["data"]
        p = sum(1 for r in d if r[8] == "PASS")
        k = sum(1 for r in d if r[8] == "SKIP")
        f = sum(1 for r in d if r[8] == "FAIL")
        print(f"  {s['name']:<23}  {len(d):>4}  {p:>4}  {k:>4}  {f:>4}")
        total += len(d)
    print("─" * 50)
    print(f"  {'TOTAL':<23}  {total:>4}")


if __name__ == "__main__":
    main()
