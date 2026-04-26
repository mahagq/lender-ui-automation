"""
ApplicationDataReader — Robot Framework keyword library
Reads application_data.xlsx and returns a list of row dictionaries.
"""
from pathlib import Path
from openpyxl import load_workbook


class ApplicationDataReader:
    ROBOT_LIBRARY_SCOPE = "GLOBAL"

    def read_application_data(self, file_path: str) -> list:
        """
        Read *file_path* (XLSX) and return a list of dicts, one per data row.
        The first row must be the header.  Empty rows are skipped.
        """
        wb = load_workbook(filename=str(Path(file_path)), read_only=True, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        result = []

        for row in rows[1:]:
            # Skip completely empty rows
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            record = {}
            for header, cell in zip(headers, row):
                record[header] = "" if cell is None else str(cell).strip()
            result.append(record)

        wb.close()
        return result
